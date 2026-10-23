"""PS-based results reporting. No causal claims inferred from activity counts.
Each KPI observation retains its own unit, period, baseline, target and method.
"""
import csv,io,json
from pathlib import Path
LANGS=['fr','en','rw','sw']
WORDS={
 'title':['Rapport axé sur les résultats du PS','PS results-based report','Raporo ishingiye ku bisubizo bya PS','Ripoti inayozingatia matokeo ya PS'],
 'draft':['BROUILLON — validation MEL, vérification Coordinateur, approbation Directeur requises','DRAFT — MEL validation, Coordinator verification and Director approval required','IGITEGANYO — gikeneye MEL, Coordinator na Director','RASIMU — inahitaji uthibitishaji wa MEL, uhakiki wa Mratibu na idhini ya Mkurugenzi'],
 'summary':['Synthèse des résultats','Results summary','Incamake y’ibyagezweho','Muhtasari wa matokeo'],
 'outcomes':['Analyse des changements et contribution','Analysis of changes and contribution','Isesengura ry’impinduka n’umusanzu','Uchambuzi wa mabadiliko na mchango'],
 'variance':['Écarts, causes et hypothèses','Variances, causes and assumptions','Itandukaniro, impamvu n’ibyatekerejwe','Tofauti, sababu na dhana'],
 'learning':['Apprentissages et adaptations','Learning and adaptations','Amasomo n’impinduka mu mikorere','Mafunzo na marekebisho'],
 'actions':['Décisions, responsables et échéances','Decisions, owners and deadlines','Ibyemezo, ababishinzwe n’igihe','Maamuzi, wahusika na muda'],
 'limitations':['Qualité des preuves et limites','Evidence quality and limitations','Ubwiza bw’ibimenyetso n’imbogamizi','Ubora wa ushahidi na mipaka'],
 'result':['Résultat du PS','PS result','Igisubizo cya PS','Matokeo ya PS'],
 'matrix':['Matrice des résultats et indicateurs','Results and indicators matrix','Imbonerahamwe y’ibisubizo n’ibipimo','Jedwali la matokeo na viashiria'],
 'indicator':['Indicateur','Indicator','Igipimo','Kiashiria'],
 'baseline':['Référence','Baseline','Agaciro fatizo','Thamani ya mwanzo'],
 'target':['Cible de référence','Reference target','Intego y’igereranya','Lengo la rejea'],
 'actual':['Valeur observée','Observed value','Agaciro kapimwe','Thamani iliyoonekana'],
 'change':['Évolution depuis la référence','Change from baseline','Impinduka kuva ku gaciro fatizo','Mabadiliko kutoka mwanzo'],
 'evidence':['Preuve / méthode','Evidence / method','Ikimenyetso / uburyo','Ushahidi / mbinu'],
 'na':['Non disponible','Unavailable','Ntibiboneka','Haipatikani'],
 'outputs':['Produits réalisés et prévus','Actual and planned outputs','Umusaruro wagezweho n’uwateganyijwe','Mazao yaliyopatikana na yaliyopangwa'],
 'finance':['Ressources, efficacité et efficience','Resources, effectiveness and efficiency','Umutungo, kugera ku ntego n’imikoreshereze','Rasilimali, ufanisi na matumizi'],
 'unmapped':['Observations sans rattachement explicite à un résultat PS','Observations without an explicit PS result link','Ibipimo bitarahuzwa n’igisubizo cya PS','Vipimo visivyounganishwa wazi na matokeo ya PS'],
 'coverage':['Couverture, portée et provenance','Coverage, scope and provenance','Aho amakuru ava n’inzego areba','Wigo na chanzo'],
 'note':['Les produits ne prouvent pas un effet ou un impact. Aucune causalité ni atteinte du résultat stratégique n’est déduite automatiquement. Les observations de périmètres différents ne sont ni additionnées ni moyennées. Les cibles annuelles ne sont pas transformées en cibles trimestrielles.','Outputs do not establish outcomes or impact. No causality or strategic result achievement is inferred automatically. Observations from different scopes are not summed or averaged. Annual targets are not converted into quarterly targets.','Umusaruro w’ibikorwa si gihamya y’impinduka cyangwa ingaruka. Nta sano y’impamvu n’ingaruka ihita yemezwa. Ibipimo byo mu nzego zitandukanye ntibiteranywa. Intego y’umwaka ntihindurwa iy’igihembwe.','Mazao ya shughuli hayathibitishi matokeo au athari. Uhusiano wa sababu na athari hauamuliwi kiotomatiki. Vipimo vya maeneo tofauti havijumlishwi. Malengo ya mwaka hayageuzwi kuwa ya robo mwaka.']}
WORDS.update({
 'axis':['Axe','Axis','Axe','Mhimili'],'objective':['Objectif','Objective','Intego','Lengo'],
 'unit':['Unité','Unit','Urwego / Igipimo','Kitengo / Kipimo'],'date':['Date','Date','Itariki','Tarehe'],
 'variance_number':['Écart à la cible','Variance from target','Itandukaniro ku ntego','Tofauti na lengo'],
 'method':['Méthode','Method','Uburyo','Mbinu'],'budget':['Budget RWF','Budget RWF','Ingengo RWF','Bajeti RWF'],
 'spent':['Dépenses RWF','Expenditure RWF','Iyakoreshejwe RWF','Matumizi RWF'],
 'owner':['Responsable','Owner','Ushinzwe','Mhusika'],'due':['Échéance','Deadline','Igihe ntarengwa','Tarehe ya mwisho'],
 'effect':['Effet vérifié','Verified effect','Impinduka yagenzuwe','Matokeo yaliyohakikiwa'],
 'lesson':['Leçon','Lesson','Isomo','Funzo'],'action':['Action','Action','Igikorwa','Hatua'],
 'source':['Source','Source','Inkomoko','Chanzo'],'revision':['Révision','Revision','Verisiyo','Toleo']})
NARRATIVE=['summary','outcomes','variance','learning','actions','limitations']
def label(x,lang):return x.get('labels',x.get('label',{})).get(lang) or x.get('labels',x.get('label',{})).get('fr','')
def words(lang):return {k:v[LANGS.index(lang)]for k,v in WORDS.items()}
def narrative_checked(n):
 from engine import Fault
 if n is None:return {}
 if not isinstance(n,dict) or set(n)-set(NARRATIVE):raise Fault(422,'invalid_results_narrative')
 if any(not isinstance(v,str) or len(v)>3000 for v in n.values()):raise Fault(422,'results_narrative_limit')
 return {k:v.strip() for k,v in n.items()}
def model(rows,config,start,end,lang,narrative=None):
 from engine import ROOT
 narrative=narrative_checked(narrative);approved=[r for r in rows if r['state'] in ['APPROVED','LOCKED'] and r['form'] in ['F03','F07','F19','F29','F30','F42']];known={x['code']:x for x in json.loads((ROOT/'registers/kpi_dictionary.json').read_text(encoding="utf-8"))};results=config['strategicResults'];valid={x['code']for x in results};observations=[]
 for r in approved:
  if r['form']!='F42':continue
  v=r['values'];k=known.get(v.get('indicator'))
  if not k:continue
  n=float(v['numerator']);d=float(v['denominator'])if v.get('denominator')not in [None,''] else None;kind=k['kind'];actual=None if kind=='gap' else n if kind in ['count','median'] else None if d in [None,0] else n/d*(1 if kind=='unit' else 100)
  baseline=float(v['baseline'])if v.get('baseline')not in [None,''] else None;target=float(v['target'])if v.get('target')not in [None,''] else None
  observations.append(dict(axis=v.get('axis') or v['indicator'].split('-')[0],result=v.get('result')if v.get('result')in valid and v['result'].split('.')[0]==v['indicator'].split('-')[0] else None,indicator=v['indicator'],label=label(k,lang),kind=kind,baseline=baseline,target=target,actual=actual,change=None if actual is None or baseline is None else actual-baseline,variance=None if actual is None or target is None else actual-target,unit=r['unit'],unitLabel=next((label(x,lang) for x in config.get('organisationRegister',[]) if x.get('dhis2_uid')==r['unit']),r['unit']),intervention=v.get('intervention'),ap=v.get('ap'),project=v.get('project'),date=r['event_date'],evidence=v.get('evidence'),method=v.get('method'),source=r['id'],revision=r['revision']))
 # Only explicit PS mappings enter the result matrix. No result is inferred from an axis.
 sections=[]
 for axis in config['options']['axis']:
  objective=next((o for o in config['objectives']if o['axis']==axis['code']),None)
  sections.append(dict(axis=axis['code'],label=label(axis,lang),objective=label(objective,lang)if objective else '',results=[dict(code=x['code'],label=label(x,lang),observations=[o for o in observations if o['result']==x['code']])for x in results if x['code'].startswith(axis['code']+'.')]))
 plans=[r for r in approved if r['form']=='F03'];outputs=[]
 for p in plans:
  v=p['values'];actuals=[r for r in approved if r['form']=='F07' and r['values'].get('plan_ref')==p['id'] and v.get('output_unit') and r['values'].get('output_unit')==v['output_unit'] and r['values'].get('outputs')not in [None,'']];expenses=[r for r in approved if r['form']=='F19' and r['values'].get('plan_ref')==p['id'] and r['values'].get('spent')not in [None,'']]
  actual=sum(float(r['values']['outputs'])for r in actuals)if actuals else None;target=float(v['target'])if v.get('target')not in [None,'']else None;budget=float(v['budget'])if v.get('budget')not in [None,'']else None;spent=sum(float(r['values']['spent'])for r in expenses)if expenses else None
  outputs.append(dict(axis=v.get('axis'),result=v.get('result'),objective=v.get('objective'),intervention=v.get('intervention'),ap=v.get('ap'),plan=p['id'],title=v.get('title'),project=v.get('project'),work_type=v.get('work_type'),partners=v.get('partners'),unit=v.get('output_unit'),target=target,actual=actual,referenceStart=v.get('start'),referenceEnd=v.get('end'),budget=budget,spent=spent,effectiveness=None if actual is None or not target else actual/target*100,costPerOutput=None if spent is None or not actual else spent/actual,evidence=[r['values'].get('evidence')for r in actuals]))
 learning=[dict(form=r['form'],source=r['id'],lesson=r['values'].get('lesson'),action=r['values'].get('action'),owner=r['values'].get('owner'),due=r['values'].get('due'),effect=(r.get('completion')or {}).get('values',{}).get('effect'))for r in approved if r['form'] in ['F29','F30']]
 return dict(strategySource=config.get('reporting',{}).get('basis','PS 2026–2031 final 03.09.2026'),code='R13',language=lang,start=start,end=end,sections=sections,unmapped=[o for o in observations if not o['result']],outputs=outputs,learning=learning,narrative=narrative,provenance=[dict(id=r['id'],revision=r['revision'],form=r['form'],unit=r['unit'])for r in approved if r['form'] in ['F03','F07','F29','F30','F42']],approvedSourceRecords=len(approved),status='DRAFT')
def render(data,lang,fmt):
 from reports import safe
 t=words(lang);flat=[]
 for a in data['sections']:
  for r in a['results']:
   for o in r['observations'] or [{}]:flat.append([a['axis'],a['objective'],r['code']+' '+r['label'],o.get('indicator'),o.get('unitLabel',o.get('unit')),o.get('date'),o.get('baseline'),o.get('target'),o.get('actual'),o.get('change'),o.get('variance'),o.get('evidence'),o.get('method'),o.get('source'),o.get('revision')])
 headers=[t['axis'],t['objective'],t['result'],t['indicator'],t['unit'],t['date'],t['baseline'],t['target'],t['actual'],t['change'],t['variance_number'],t['evidence'],t['method'],t['source'],t['revision']]
 if fmt=='csv':
  out=io.StringIO();w=csv.writer(out);w.writerow(headers)
  for row in flat:w.writerow([safe(v)if v is not None else t['na']for v in row])
  return ('\ufeff'+out.getvalue()).encode(),'text/csv;charset=utf-8'
 if fmt=='xlsx':
  from openpyxl import Workbook
  from openpyxl.styles import Font,PatternFill,Alignment
  wb=Workbook();ws=wb.active;ws.title='PS';ws.append(headers)
  for row in flat:ws.append([safe(v)if isinstance(v,str)else v for v in row])
  for name,rows in [('Narrative',[[t[k],data['narrative'].get(k,t['na'])]for k in NARRATIVE]),('Outputs',[[t['axis'],t['result'],'Intervention','AP','Projet / Project',t['unit'],t['target'],t['actual'],t['budget'],t['spent'],'EFF-01 %','EFI-01 RWF',t['source']]]+[[x[k]for k in ['axis','result','intervention','ap','project','unit','target','actual','budget','spent','effectiveness','costPerOutput','plan']]for x in data['outputs']]),('Unmapped',[[o['indicator'],o['unit'],o['actual'],o['evidence']]for o in data['unmapped']]),('Learning',[[t['source'],t['lesson'],t['action'],t['owner'],t['due'],t['effect']]]+[[x['source'],x['lesson'],x['action'],x['owner'],x['due'],x['effect']]for x in data['learning']]),('Sources',[[x['id'],x['revision'],x['form'],x['unit']]for x in data['provenance']])]:
   sheet=wb.create_sheet(name)
   for row in rows:sheet.append([safe(v)if isinstance(v,str)else v for v in row])
  for sheet in wb:
   sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
   for cell in sheet[1]:cell.font=Font(name='Arial',bold=True,size=12);cell.fill=PatternFill('solid',fgColor='FDF751')
   for col in sheet.columns:
    sheet.column_dimensions[col[0].column_letter].width=28
    for cell in col:cell.alignment=Alignment(wrap_text=True,vertical='top')
  out=io.BytesIO();wb.save(out);return out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
 from docx import Document
 from docx.shared import Pt,Inches
 doc=Document();sec=doc.sections[0];sec.left_margin=sec.right_margin=Inches(.65)
 for style,size in [('Normal',12),('Title',24),('Heading 1',18),('Heading 2',12)]:doc.styles[style].font.name='Arial';doc.styles[style].font.size=Pt(size)
 doc.add_heading(t['title'],0);doc.add_paragraph(data['strategySource']);doc.add_paragraph(data['start']+' — '+data['end']);doc.add_paragraph(t['draft']);doc.add_paragraph(t['note'])
 for k in ['summary','outcomes']:doc.add_heading(t[k],1);doc.add_paragraph(data['narrative'].get(k)or t['na'])
 for a in data['sections']:
  doc.add_heading(a['axis']+' · '+a['label'],1);doc.add_paragraph(a['objective'])
  for r in a['results']:
   doc.add_heading(r['code']+' · '+r['label'],2)
   if not r['observations']:doc.add_paragraph(t['na']);continue
   table=doc.add_table(rows=1,cols=4);table.style='Table Grid'
   for cell,value in zip(table.rows[0].cells,[t['indicator'],t['baseline']+' / '+t['target'],t['actual']+' / '+t['change'],t['evidence']]):cell.text=value
   for o in r['observations']:
    row=table.add_row().cells
    for cell,value in zip(row,[o['indicator']+' '+o['label']+'\n'+o.get('unitLabel',o['unit'])+' · '+o['date'],str(o['baseline'])+' / '+str(o['target']),str(o['actual'])+' / '+str(o['change']),(o['evidence']or'')+'\n'+(o['method']or'')]):cell.text=value.replace('None',t['na'])
 doc.add_heading(t['outputs'],1)
 for x in data['outputs']:
  doc.add_heading((x['result']or'')+' · '+(x['title']or''),2);doc.add_paragraph(' / '.join(str(x.get(k)or t['na'])for k in ['objective','intervention','ap','project','work_type']))
  doc.add_paragraph(str(x['referenceStart'])+' — '+str(x['referenceEnd']))
  table=doc.add_table(rows=1,cols=4);table.style='Table Grid'
  for cell,value in zip(table.rows[0].cells,[t['target'],t['actual'],'EFF-01 %',t['unit']]):cell.text=value
  for cell,value in zip(table.add_row().cells,[x['target'],x['actual'],x['effectiveness'],x['unit']]):cell.text=t['na']if value is None else str(value)
  doc.add_heading(t['finance'],2);table=doc.add_table(rows=1,cols=3);table.style='Table Grid'
  for cell,value in zip(table.rows[0].cells,[t['budget'],t['spent'],'EFI-01 RWF']):cell.text=value
  for cell,value in zip(table.add_row().cells,[x['budget'],x['spent'],x['costPerOutput']]):cell.text=t['na']if value is None else str(value)
  doc.add_paragraph(t['evidence']+': '+', '.join(v for v in x['evidence']if v))
 doc.add_heading(t['learning'],1)
 for x in data['learning']:
  doc.add_paragraph((x.get('lesson')or'')+' — '+(x.get('action')or''));doc.add_paragraph(' · '.join(str(x.get(k)or t['na'])for k in ['owner','due','effect']));doc.add_paragraph(x['source'])
 for k in ['variance','learning','actions','limitations']:doc.add_heading(t[k],1);doc.add_paragraph(data['narrative'].get(k)or t['na'])
 doc.add_heading(t['unmapped'],1)
 for o in data['unmapped']:doc.add_paragraph(o['indicator']+' · '+o['unit']+' · '+str(o['actual']))
 doc.add_heading(t['coverage'],1);doc.add_paragraph(str(data['approvedSourceRecords'])+' · '+data['start']+' — '+data['end'])
 for x in data['provenance']:doc.add_paragraph(x['id']+' · r'+str(x['revision'])+' · '+x['form'])
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 for table in doc.tables:
  for cell in table.rows[0].cells:
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FDF751');cell._tc.get_or_add_tcPr().append(shade)
   for paragraph in cell.paragraphs:
    for run in paragraph.runs:run.bold=True
 out=io.BytesIO();doc.save(out);return out.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
