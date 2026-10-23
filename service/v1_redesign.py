"""APSE v1.0 Redesign Backend: Task-driven To-Do workflow, automatic metadata inheritance,
location auto-scoping, 4-step guided capture, impact threshold engine, 3-tier approvals,
outcome accordions, regional parish map, and plain-language validation.
"""
import csv, hashlib, io, json, math, re, uuid
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

# Natural plain-language translations for Strategic Axes & Outcomes (EN, RW, FR, SW)
AXIS_NATURAL = {
    'A1': {
        'en': 'Healing, Unity & Reconciliation',
        'rw': 'Isanamitima, Ubumwe n’Ubwiyunge',
        'fr': 'Guérison, Unité et Réconciliation',
        'sw': 'Uponyaji, Umoja na Upatanishishi',
    },
    'A2': {
        'en': 'Peaceful Conflict Resolution',
        'rw': 'Gukemura Amakimbirane mu Mahoro',
        'fr': 'Résolution Pacifique des Conflits',
        'sw': 'Utatuzi wa Migogoro kwa Amani',
    },
    'A3': {
        'en': 'Environmental Care & Climate Action',
        'rw': 'Kubungabunga Ibidukikije',
        'fr': 'Justice Écologique et Environnement',
        'sw': 'Utunzaji wa Mazingira',
    },
    'A4': {
        'en': 'Human Rights & Legal Aid',
        'rw': 'Uburenganzira bwa Muntu n’Ubuvugizi',
        'fr': 'Droits Humains et Aide Juridique',
        'sw': 'Haki za Binadamu na Msaada wa Kisheria',
    },
    'A5': {
        'en': 'Parish Committee & Team Strengthening',
        'rw': 'Kubaka Ubushobozi bwa Komite n’Abakozi',
        'fr': 'Renforcement des Comités et Équipes',
        'sw': 'Kuimarisha Kamati na Timu',
    },
}

ISSUE_CATEGORIES = ['Transport', 'Turnout', 'Budget', 'Weather', 'Security', 'Other']

PARISH_DISTRICTS = {
    'PAR-01': ('Gakoma', 'Gisagara', 68, 38),
    'PAR-02': ('Magi', 'Gisagara', 62, 52),
    'PAR-03': ('Save', 'Gisagara', 58, 44),
    'PAR-04': ('Kansi', 'Gisagara', 64, 68),
    'PAR-05': ('Kirarambogo', 'Gisagara', 74, 62),
    'PAR-06': ('Mugombwa', 'Gisagara', 78, 72),
    'PAR-07': ('Ndora', 'Gisagara', 70, 54),
    'PAR-08': ('Gishubi', 'Gisagara', 66, 46),
    'PAR-09': ('Higiro', 'Gisagara', 60, 60),
    'PAR-10': ('Butare (Cathedral)', 'Huye', 48, 50),
    'PAR-11': ('Ngoma', 'Huye', 45, 55),
    'PAR-12': ('Tumba', 'Huye', 44, 60),
    'PAR-13': ('Mbazi', 'Huye', 46, 42),
    'PAR-14': ('Ruhashya', 'Huye', 40, 34),
    'PAR-15': ('Rusatira', 'Huye', 43, 28),
    'PAR-16': ('Songa', 'Huye', 50, 26),
    'PAR-17': ('Karama', 'Huye', 38, 48),
    'PAR-18': ('Simbi', 'Huye', 35, 56),
    'PAR-19': ('Maraba', 'Huye', 36, 42),
    'PAR-20': ('Nyanza', 'Nyanza', 48, 18),
    'PAR-21': ('Muyunzwe', 'Nyanza', 42, 14),
    'PAR-22': ('Gatagara', 'Nyanza', 54, 20),
    'PAR-23': ('Nyumba', 'Huye', 52, 66),
    'PAR-24': ('Kibeho', 'Nyaruguru', 24, 64),
    'PAR-25': ('Cyahinda', 'Nyaruguru', 30, 76),
    'PAR-26': ('Muganza', 'Nyaruguru', 22, 72),
    'PAR-27': ('Ruramba', 'Nyaruguru', 20, 54),
}


def compute_pastoral(event_date_str):
    try:
        d = date.fromisoformat(str(event_date_str))
    except (ValueError, TypeError):
        return None
    start_year = d.year if d.month >= 9 else d.year - 1
    if start_year < 2026 or start_year > 2030:
        return None
    ap_index = start_year - 2025
    quarter = 1 if d.month in (9, 10, 11) else 2 if d.month in (12, 1, 2) else 3 if d.month in (3, 4, 5) else 4
    return {
        'ap': f'AP{ap_index}',
        'year_label': f'{start_year}–{start_year + 1}',
        'quarter': f'Q{quarter}',
        'quarter_num': quarter,
        'start': f'{start_year}-09-01',
        'end': f'{start_year + 1}-08-31',
    }


def compute_impact_rating(outcome_value, min_threshold, target_threshold):
    if outcome_value is None:
        return 'PENDING'
    val = float(outcome_value)
    min_t = float(min_threshold or 0)
    targ_t = float(target_threshold or max(min_t, 1))
    if val >= targ_t:
        return 'GREAT'
    if val >= min_t:
        return 'GOOD'
    return 'LOW'


class V1Redesign:
    def init_v1(self):
        self.db.executescript('''
CREATE TABLE IF NOT EXISTS v1_tasks(
    id TEXT PRIMARY KEY,
    code TEXT NOT NULL,
    unit TEXT NOT NULL,
    parish_code TEXT NOT NULL,
    axis TEXT NOT NULL,
    outcome TEXT NOT NULL,
    intervention TEXT NOT NULL,
    indicator TEXT NOT NULL,
    ap TEXT NOT NULL,
    pastoral_quarter TEXT NOT NULL,
    due_date TEXT NOT NULL,
    planned_target REAL NOT NULL,
    min_threshold REAL NOT NULL,
    target_threshold REAL NOT NULL,
    frequency INTEGER NOT NULL DEFAULT 1,
    is_sensitive INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'PENDING',
    created_by TEXT NOT NULL,
    body BLOB NOT NULL,
    created TEXT NOT NULL,
    updated TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS v1_logs(
    id TEXT PRIMARY KEY,
    task_id TEXT NOT NULL,
    unit TEXT NOT NULL,
    parish_code TEXT NOT NULL,
    centrale TEXT NOT NULL,
    event_date TEXT NOT NULL,
    ap TEXT NOT NULL,
    pastoral_quarter TEXT NOT NULL,
    collection_mode TEXT NOT NULL,
    women INTEGER NOT NULL DEFAULT 0,
    men INTEGER NOT NULL DEFAULT 0,
    total INTEGER NOT NULL DEFAULT 0,
    age_u18 INTEGER NOT NULL DEFAULT 0,
    age_18_35 INTEGER NOT NULL DEFAULT 0,
    age_36_64 INTEGER NOT NULL DEFAULT 0,
    age_65 INTEGER NOT NULL DEFAULT 0,
    disability INTEGER NOT NULL DEFAULT 0,
    sessions_done REAL NOT NULL DEFAULT 1,
    outcome_value REAL NOT NULL DEFAULT 0,
    impact_rating TEXT NOT NULL DEFAULT 'GOOD',
    issue_severity TEXT NOT NULL DEFAULT 'NONE',
    issue_categories TEXT NOT NULL DEFAULT '',
    issue_note TEXT NOT NULL DEFAULT '',
    stage TEXT NOT NULL DEFAULT 'SUBMITTED',
    revision INTEGER NOT NULL DEFAULT 1,
    version INTEGER NOT NULL DEFAULT 1,
    parent_id TEXT,
    author TEXT NOT NULL,
    body BLOB NOT NULL,
    created TEXT NOT NULL,
    updated TEXT NOT NULL
);
''')
        self._seed_v1_if_empty()

    def _seed_v1_if_empty(self):
        from engine import NOW
        with self.lock:
            count = self.db.execute('SELECT COUNT(*) FROM v1_tasks').fetchone()[0]
            if count > 0:
                return
            stamp = NOW()
            seed_tasks = [
                {
                    'id': 'task-a2-gakoma-01',
                    'code': 'TSK-2026-001',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-01',
                    'axis': 'A2',
                    'outcome': 'A2.R1',
                    'intervention': 'A2.I1',
                    'indicator': 'KPI-A2-01',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'due_date': '2026-12-15',
                    'planned_target': 20,
                    'min_threshold': 12,
                    'target_threshold': 18,
                    'frequency': 4,
                    'is_sensitive': 0,
                    'status': 'IN_PROGRESS',
                    'titles': {
                        'en': 'Train 20 Centrale Mediators in Early Family Conflict Resolution',
                        'rw': 'Guhugura Abahuza 20 ba Santarali mu Gukemura Amakimbirane mu Miryango',
                        'fr': 'Former 20 médiateurs de Centrale à la résolution précoce des conflits familiaux',
                        'sw': 'Kuwafunza Wapatanishi 20 wa Kigango katika Utatuzi wa Migogoro ya Familia',
                    },
                    'outcome_prompt': {
                        'en': 'How many trained mediators actively resolved a family dispute this period?',
                        'rw': 'Ni abahuza bangahe bahuguwe bakemuye amakimbirane mu muryango muri iki gihe?',
                        'fr': 'Combien de médiateurs formés ont résolu un conflit familial sur la période ?',
                        'sw': 'Wapatanishi wangapi waliofunzwa walitatua mgogoro wa familia kipindi hiki?',
                    },
                    'unit_label': 'mediators trained',
                },
                {
                    'id': 'task-a1-butare-02',
                    'code': 'TSK-2026-002',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-10',
                    'axis': 'A1',
                    'outcome': 'A1.R1',
                    'intervention': 'A1.I1',
                    'indicator': 'KPI-A1-01',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'due_date': '2026-11-30',
                    'planned_target': 4,
                    'min_threshold': 30,
                    'target_threshold': 45,
                    'frequency': 4,
                    'is_sensitive': 0,
                    'status': 'COMPLETED',
                    'titles': {
                        'en': 'Facilitate Isanamitima & GURR Healing Circles (Q1)',
                        'rw': 'ikiganiro cya Isanamitima na GURR mu matsinda (Igihembwe cya 1)',
                        'fr': 'Animer les cercles de guérison Isanamitima et GURR (T1)',
                        'sw': 'Kuendesha miduara ya uponyaji ya Isanamitima na GURR (Robo ya 1)',
                    },
                    'outcome_prompt': {
                        'en': 'Participants reporting improved trust and emotional safety (post-session check)',
                        'rw': 'abitabiriye bagaragaje ko icyizere n’ituze ry’umutima byiyongereye',
                        'fr': 'Participants déclarant une amélioration de la confiance et de la sécurité émotionnelle',
                        'sw': 'Washiriki walioripoti kuongezeka kwa imani na utulivu wa kihisia',
                    },
                    'unit_label': 'sessions held',
                },
                {
                    'id': 'task-a4-save-03',
                    'code': 'TSK-2026-003',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-03',
                    'axis': 'A4',
                    'outcome': 'A4.R1',
                    'intervention': 'A4.I1',
                    'indicator': 'KPI-A4-01',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'due_date': '2026-11-20',
                    'planned_target': 150,
                    'min_threshold': 90,
                    'target_threshold': 130,
                    'frequency': 2,
                    'is_sensitive': 0,
                    'status': 'PENDING',
                    'titles': {
                        'en': 'Community Legal Rights & Land Succession Awareness Session',
                        'rw': 'Ubukangurambaga ku Burenganzira mu Mategeko n’Izungura ry’Ubutaka',
                        'fr': 'Sensibilisation communautaire sur les droits fonciers et successoraux',
                        'sw': 'Elimu ya Jamii kuhusu Haki za Kisheria na Urithi wa Ardhi',
                    },
                    'outcome_prompt': {
                        'en': 'Citizens passing the post-session rights knowledge check or referred for legal aid',
                        'rw': 'abaturage bamenye neza amategeko y’izungura cyangwa bahawe ubufasha mu mategeko',
                        'fr': 'Citoyens ayant réussi le test de connaissances ou orientés vers l’aide juridique',
                        'sw': 'Wananchi waliofaulu jaribio la haki au kupewa rufaa ya msaada wa kisheria',
                    },
                    'unit_label': 'citizens reached',
                },
                {
                    'id': 'task-a3-kibeho-04',
                    'code': 'TSK-2026-004',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-24',
                    'axis': 'A3',
                    'outcome': 'A3.R1',
                    'intervention': 'A3.I1',
                    'indicator': 'KPI-A3-01',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'due_date': '2026-12-05',
                    'planned_target': 500,
                    'min_threshold': 300,
                    'target_threshold': 450,
                    'frequency': 2,
                    'is_sensitive': 0,
                    'status': 'IN_PROGRESS',
                    'titles': {
                        'en': 'Youth Laudato Si’ Agroforestry & Soil Protection Action',
                        'rw': 'Igikorwa cy’Urubyiruko cya Laudato Si’ cyo Gutera Ibiti no Kurwanya Isuri',
                        'fr': 'Action jeunesse Laudato Si’ d’agroforesterie et protection des sols',
                        'sw': 'Hatua ya Vijana ya Laudato Si’ ya Kupanda Miti na Kuhifadhi Udongo',
                    },
                    'outcome_prompt': {
                        'en': 'Surviving agroforestry trees verified & households adopting soil protection',
                        'rw': 'ibiti byatewe byafashe n’ingo zashyizeho uburyo bwo kurwanya isuri',
                        'fr': 'Arbres agroforestiers vérifiés et ménages adoptant la protection des sols',
                        'sw': 'Miti iliyostawi na kaya zilizokubali mbinu za kuhifadhi udongo',
                    },
                    'unit_label': 'trees / households',
                },
                {
                    'id': 'task-a1-ngoma-05',
                    'code': 'TSK-2026-005',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-11',
                    'axis': 'A1',
                    'outcome': 'A1.R3',
                    'intervention': 'A1.I3',
                    'indicator': 'KPI-A1-03',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'due_date': '2026-12-20',
                    'planned_target': 15,
                    'min_threshold': 10,
                    'target_threshold': 14,
                    'frequency': 4,
                    'is_sensitive': 1,
                    'status': 'PENDING',
                    'titles': {
                        'en': 'Log Confidential Case: Prison Reintegration & Listening Accompaniment',
                        'rw': 'Andika Ikibazo cy’Ibanga: Gukurikirana Abagororwa n’Abasubira mu Muryango',
                        'fr': 'Enregistrer un cas confidentiel : écoute et réintégration post-carcérale',
                        'sw': 'Rekodi Kesi ya Siri: Usikilizaji na Ujumuishaji Upya baada ya Gereza',
                    },
                    'outcome_prompt': {
                        'en': 'Accompanied individuals safely reintegrated with family acceptance',
                        'rw': 'abafashijwe basubiye mu miryango yabo mu mahoro no kwakirwa neza',
                        'fr': 'Personnes accompagnées réintégrées en sécurité avec acceptation familiale',
                        'sw': 'Watu waliosaidiwa kujiunga tena na familia zao kwa amani',
                    },
                    'unit_label': 'cases accompanied',
                },
                {
                    'id': 'task-a5-nyanza-06',
                    'code': 'TSK-2026-006',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-20',
                    'axis': 'A5',
                    'outcome': 'A5.R1',
                    'intervention': 'A5.I1',
                    'indicator': 'KPI-A5-01',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'due_date': '2026-12-10',
                    'planned_target': 6,
                    'min_threshold': 4,
                    'target_threshold': 6,
                    'frequency': 4,
                    'is_sensitive': 0,
                    'status': 'PENDING',
                    'titles': {
                        'en': 'Quarterly CPJP Parish Committee Governance & Data Quality Review',
                        'rw': 'Inama y’Igihembwe ya Komite CPJP yo Gusuzuma Imikorere n’Amakuru',
                        'fr': 'Revue trimestrielle de gouvernance et qualité des données du comité CPJP',
                        'sw': 'Mapitio ya Robo Mwaka ya Utawala na Ubora wa Data ya Kamati ya CPJP',
                    },
                    'outcome_prompt': {
                        'en': 'Centrale committees submitting complete, verified quarterly activity books',
                        'rw': 'komite za santarali zatanze raporo zuzuye kandi zagenzuwe',
                        'fr': 'Comités de centrale soumettant des cahiers d’activités complets et vérifiés',
                        'sw': 'Kamati za vigango zilizowasilisha vitabu kamili vya shughuli',
                    },
                    'unit_label': 'centrales verified',
                },
            ]
            for t in seed_tasks:
                body = {
                    'titles': t['titles'],
                    'outcome_prompt': t['outcome_prompt'],
                    'unit_label': t['unit_label'],
                }
                self.db.execute(
                    'INSERT INTO v1_tasks VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (
                        t['id'], t['code'], t['unit'], t['parish_code'], t['axis'],
                        t['outcome'], t['intervention'], t['indicator'], t['ap'],
                        t['pastoral_quarter'], t['due_date'], t['planned_target'],
                        t['min_threshold'], t['target_threshold'], t['frequency'],
                        t['is_sensitive'], t['status'], 'SYSTEM_SEED',
                        self.enc(body), stamp, stamp,
                    ),
                )

            seed_logs = [
                {
                    'id': 'log-seed-001',
                    'task_id': 'task-a1-butare-02',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-10',
                    'centrale': 'Ste Thérèse',
                    'event_date': '2026-10-10',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'collection_mode': 'FOCUS_GROUP',
                    'women': 28,
                    'men': 20,
                    'total': 48,
                    'age_u18': 4,
                    'age_18_35': 18,
                    'age_36_64': 20,
                    'age_65': 6,
                    'disability': 3,
                    'sessions_done': 4,
                    'outcome_value': 46,
                    'impact_rating': 'GREAT',
                    'issue_severity': 'NONE',
                    'issue_categories': '',
                    'issue_note': '',
                    'stage': 'DIOCESE_APPROVED',
                    'author': 'CPJP_BUTARE',
                    'payload': {
                        'duration_hours': 3,
                        'consensus_notes': 'Strong trust built between 14 families; 3 longstanding neighbour disputes reconciled.',
                        'story_of_change': 'Two families in Rango cell restored shared boundary farming after 4 sessions.',
                        'attachments': ['Attendance_Sheet_Butare_Oct2026.pdf'],
                        'privacy_confirmed': True,
                        'reviews': [
                            {'stage': 'PARISH_APPROVED', 'actor': 'CPJP Committee Butare', 'time': '2026-10-11T10:00:00Z', 'note': 'Verified attendance list and facilitator notes.'},
                            {'stage': 'DIOCESE_APPROVED', 'actor': 'CDJP Diocese MEL', 'time': '2026-10-12T14:30:00Z', 'note': 'Approved and included in Q1 Diocesan Impact totals.'},
                        ],
                    },
                },
                {
                    'id': 'log-seed-002',
                    'task_id': 'task-a2-gakoma-01',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-01',
                    'centrale': 'Kizenga',
                    'event_date': '2026-10-18',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'collection_mode': 'ENUMERATOR',
                    'women': 7,
                    'men': 5,
                    'total': 12,
                    'age_u18': 0,
                    'age_18_35': 5,
                    'age_36_64': 6,
                    'age_65': 1,
                    'disability': 1,
                    'sessions_done': 12,
                    'outcome_value': 10,
                    'impact_rating': 'LOW',
                    'issue_severity': 'MEDIUM',
                    'issue_categories': 'Budget,Transport',
                    'issue_note': 'Disbursement of transport allowance for distant CEBs arrived 1 week late, reducing turnout from 20 to 12.',
                    'stage': 'SUBMITTED',
                    'author': 'CPJP_GAKOMA',
                    'payload': {
                        'respondents': [
                            {'code': 'MED-01', 'score': 85, 'resolved_case': True},
                            {'code': 'MED-02', 'score': 90, 'resolved_case': True},
                        ],
                        'attachments': ['Training_List_Gakoma_Kizenga.pdf'],
                        'privacy_confirmed': True,
                        'reviews': [],
                    },
                },
                {
                    'id': 'log-seed-003',
                    'task_id': 'task-a3-kibeho-04',
                    'unit': 'LocalRoot01',
                    'parish_code': 'PAR-24',
                    'centrale': 'Kibeho',
                    'event_date': '2026-10-20',
                    'ap': 'AP1',
                    'pastoral_quarter': 'Q1',
                    'collection_mode': 'OBSERVATION',
                    'women': 42,
                    'men': 38,
                    'total': 80,
                    'age_u18': 22,
                    'age_18_35': 40,
                    'age_36_64': 15,
                    'age_65': 3,
                    'disability': 2,
                    'sessions_done': 340,
                    'outcome_value': 340,
                    'impact_rating': 'GOOD',
                    'issue_severity': 'MAJOR',
                    'issue_categories': 'Weather,Transport',
                    'issue_note': 'Heavy October rains washed out the access road to Mbasa hillside; remaining 160 seedlings scheduled for next week.',
                    'stage': 'PARISH_APPROVED',
                    'author': 'CPJP_KIBEHO',
                    'payload': {
                        'checklist': ['Contour trenches dug', 'Indigenous seedlings planted', 'Youth eco-club committee active'],
                        'observation_notes': '340 agroforestry seedlings planted and mulched across 3 hillside CEBs.',
                        'attachments': ['Kibeho_Agroforestry_FieldCheck.jpg'],
                        'privacy_confirmed': True,
                        'reviews': [
                            {'stage': 'PARISH_APPROVED', 'actor': 'CPJP Committee Kibeho', 'time': '2026-10-21T09:15:00Z', 'note': 'Confirmed field count with parish youth chaplain.'},
                        ],
                    },
                },
            ]
            for l in seed_logs:
                self.db.execute(
                    'INSERT INTO v1_logs VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (
                        l['id'], l['task_id'], l['unit'], l['parish_code'], l['centrale'],
                        l['event_date'], l['ap'], l['pastoral_quarter'], l['collection_mode'],
                        l['women'], l['men'], l['total'], l['age_u18'], l['age_18_35'],
                        l['age_36_64'], l['age_65'], l['disability'], l['sessions_done'],
                        l['outcome_value'], l['impact_rating'], l['issue_severity'],
                        l['issue_categories'], l['issue_note'], l['stage'], 1, 1, None,
                        l['author'], self.enc(l['payload']), stamp, stamp,
                    ),
                )

    def _parishes_and_centrales(self, user):
        m = self.metadata()
        parishes = []
        ccjp_all = m.get('ccjp_units', [])
        for p in m.get('organisation_units', []):
            if p.get('kind') != 'CPJP' or not p.get('active', True):
                continue
            pcode = p['code']
            pname, district, x, y = PARISH_DISTRICTS.get(pcode, (p['labels']['fr'], 'Huye', 50, 50))
            centrales = [
                {
                    'code': c['code'],
                    'name': c['labels'].get('rw') or c['labels'].get('fr') or c['code'],
                    'labels': c['labels'],
                }
                for c in ccjp_all
                if c.get('parent') == pcode and c.get('active', True)
            ]
            if not centrales:
                centrales = [
                    {'code': f'{pcode}-C1', 'name': f'{pname} Centre', 'labels': {'en': f'{pname} Centre', 'rw': f'{pname} Centre', 'fr': f'{pname} Centre', 'sw': f'{pname} Centre'}},
                    {'code': f'{pcode}-C2', 'name': f'Santarali ya {pname} II', 'labels': {'en': f'{pname} North', 'rw': f'Santarali ya {pname} II', 'fr': f'{pname} Nord', 'sw': f'{pname} Kaskazini'}},
                ]
            parishes.append({
                'code': pcode,
                'name': p['labels'].get('rw') or p['labels'].get('fr') or pname,
                'labels': p['labels'],
                'district': district,
                'x': x,
                'y': y,
                'centrales': centrales[:6],
                'cebs': [f'CEB {c["name"]} 1' for c in centrales[:3]] + [f'CEB {c["name"]} 2' for c in centrales[:2]],
            })
        # Determine user's default parish (auto-scoped if parish account, else PAR-01 default selection with full access)
        roles = self.roles(user)
        is_diocesan = bool(roles & {'APSE_MEL', 'APSE_COORDINATION', 'APSE_DIRECTION', 'APSE_STAFF', 'APSE_FINANCE', 'APSE_HR'})
        assigned_parish = 'PAR-01'
        for u in user.get('units', []):
            if u.get('sourceCode') and any(p['code'] == u['sourceCode'] for p in parishes):
                assigned_parish = u['sourceCode']
                break
        if not is_diocesan:
            parishes = [p for p in parishes if p['code'] == assigned_parish] or parishes[:1]
        return parishes, assigned_parish, is_diocesan

    def _goals_catalogue(self):
        m = self.metadata()
        goals = []
        for r in m.get('results', []):
            if not r.get('active', True):
                continue
            axis_code = r['axis']
            interventions = [i['code'] for i in m.get('interventions', []) if i.get('axis') == axis_code]
            goals.append({
                'outcome_code': r['code'],
                'axis_code': axis_code,
                'axis_labels': AXIS_NATURAL.get(axis_code, {'en': axis_code, 'rw': axis_code, 'fr': axis_code, 'sw': axis_code}),
                'intervention_code': interventions[0] if interventions else f'{axis_code}.I1',
                'indicator_code': f'KPI-{r["code"].replace(".", "-")}',
                'labels': r['labels'],
            })
        return goals

    def _serial_task(self, row, logs):
        body = self.dec(row['body'])
        task_logs = [l for l in logs if l['task_id'] == row['id'] and l.get('parent_id') is None]
        approved_logs = [l for l in task_logs if l['stage'] in ('PARISH_APPROVED', 'DIOCESE_APPROVED')]
        active_logs = task_logs
        completed_units = sum(float(l['sessions_done']) for l in active_logs)
        target = float(row['planned_target'] or 1)
        completion_pct = min(100.0, round((completed_units / target) * 100.0, 1)) if target > 0 else 0.0
        best_outcome = max((float(l['outcome_value']) for l in active_logs), default=None)
        impact = compute_impact_rating(best_outcome, row['min_threshold'], row['target_threshold']) if active_logs else 'PENDING'
        latest_log = active_logs[-1] if active_logs else None
        status = 'COMPLETED' if completion_pct >= 100.0 else ('IN_PROGRESS' if active_logs else row['status'])
        if latest_log and latest_log['stage'] == 'CORRECTION':
            status = 'NEEDS_CORRECTION'
        return {
            'id': row['id'],
            'code': row['code'],
            'unit': row['unit'],
            'parish_code': row['parish_code'],
            'axis': row['axis'],
            'axis_labels': AXIS_NATURAL.get(row['axis'], {}),
            'outcome': row['outcome'],
            'intervention': row['intervention'],
            'indicator': row['indicator'],
            'ap': row['ap'],
            'pastoral_quarter': row['pastoral_quarter'],
            'due_date': row['due_date'],
            'planned_target': row['planned_target'],
            'min_threshold': row['min_threshold'],
            'target_threshold': row['target_threshold'],
            'frequency': row['frequency'],
            'is_sensitive': bool(row['is_sensitive']),
            'status': status,
            'titles': body.get('titles', {}),
            'outcome_prompt': body.get('outcome_prompt', {}),
            'unit_label': body.get('unit_label', 'participants'),
            'completed_units': completed_units,
            'completion_pct': completion_pct,
            'quarters_completed': len(active_logs),
            'best_outcome': best_outcome,
            'impact_rating': impact,
            'latest_log': latest_log,
            'logs_count': len(active_logs),
            'approved_count': len(approved_logs),
        }

    def _serial_log(self, row):
        payload = self.dec(row['body'])
        return {
            'id': row['id'],
            'task_id': row['task_id'],
            'unit': row['unit'],
            'parish_code': row['parish_code'],
            'centrale': row['centrale'],
            'event_date': row['event_date'],
            'ap': row['ap'],
            'pastoral_quarter': row['pastoral_quarter'],
            'collection_mode': row['collection_mode'],
            'women': row['women'],
            'men': row['men'],
            'total': row['total'],
            'age_u18': row['age_u18'],
            'age_18_35': row['age_18_35'],
            'age_36_64': row['age_36_64'],
            'age_65': row['age_65'],
            'disability': row['disability'],
            'sessions_done': row['sessions_done'],
            'outcome_value': row['outcome_value'],
            'impact_rating': row['impact_rating'],
            'issue_severity': row['issue_severity'],
            'issue_categories': [c for c in row['issue_categories'].split(',') if c],
            'issue_note': row['issue_note'],
            'stage': row['stage'],
            'revision': row['revision'],
            'version': row['version'],
            'parent_id': row['parent_id'],
            'author': row['author'],
            'payload': payload,
            'created': row['created'],
            'updated': row['updated'],
        }

    def v1_bootstrap(self, user, parish_filter=None, period_filter=None, axis_filter=None):
        with self.lock:
            parishes, assigned_parish, is_diocesan = self._parishes_and_centrales(user)
            parish_map = {p['code']: p['name'] for p in parishes}
            raw_logs = [self._serial_log(r) for r in self.db.execute('SELECT * FROM v1_logs ORDER BY event_date, created')]
            raw_tasks = list(self.db.execute('SELECT * FROM v1_tasks ORDER BY due_date, created'))
            tasks = [self._serial_task(t, raw_logs) for t in raw_tasks]
            for t in tasks:
                t['parish_name'] = parish_map.get(t['parish_code'], t['parish_code'])
            task_lookup = {t['id']: t for t in tasks}
            for l in raw_logs:
                t = task_lookup.get(l['task_id'], {})
                l['task_titles'] = t.get('titles', {'en': 'Field Activity', 'rw': 'Igikorwa', 'fr': 'Activité', 'sw': 'Shughuli'})
                l['parish_name'] = parish_map.get(l['parish_code'], l['parish_code'])
                l['axis'] = t.get('axis', 'A1')
                l['outcome'] = t.get('outcome', 'A1.R1')
                l['planned_target'] = t.get('planned_target', 1)
                l['min_threshold'] = t.get('min_threshold', 1)
                l['target_threshold'] = t.get('target_threshold', 1)

            drafts = self.draft_listing(user, 'ACTIVE')
            analytics = self._compute_v1_analytics(tasks, raw_logs, parishes, parish_filter, period_filter, axis_filter)
            return {
                'user': {
                    'id': user['id'],
                    'name': user.get('name', 'Local User'),
                    'roles': user.get('roles', []),
                    'is_diocesan': is_diocesan,
                    'assigned_parish': assigned_parish,
                },
                'parishes': parishes,
                'goals': self._goals_catalogue(),
                'issue_categories': ISSUE_CATEGORIES,
                'tasks': tasks,
                'logs': raw_logs,
                'drafts': drafts,
                'analytics': analytics,
            }

    def v1_create_task(self, user, body):
        from engine import Fault, NOW
        title = str(body.get('title', '')).strip()
        if not title:
            raise Fault(422, 'plain_error', 'Please enter a short, clear title for this activity.')
        outcome_code = str(body.get('outcome_code', 'A1.R1')).strip()
        goals = self._goals_catalogue()
        goal = next((g for g in goals if g['outcome_code'] == outcome_code), goals[0])
        due_date = str(body.get('due_date') or '2026-12-31').strip()
        period = compute_pastoral(due_date)
        if not period:
            raise Fault(422, 'plain_error', 'Pick a date between 1 Sept 2026 and 31 Aug 2031.')
        parishes, assigned_parish, _ = self._parishes_and_centrales(user)
        parish_code = body.get('parish_code') or assigned_parish
        if not any(p['code'] == parish_code for p in parishes):
            parish_code = assigned_parish
        try:
            planned_target = float(body.get('planned_target', 20))
            min_threshold = float(body.get('min_threshold', max(1, round(planned_target * 0.6))))
            target_threshold = float(body.get('target_threshold', max(min_threshold, round(planned_target * 0.9))))
            frequency = int(body.get('frequency', 1))
        except (ValueError, TypeError):
            raise Fault(422, 'plain_error', 'Enter valid positive numbers for the target and thresholds.')
        if planned_target <= 0:
            raise Fault(422, 'plain_error', 'Planned target must be at least 1.')

        tid = 'task-' + uuid.uuid4().hex[:12]
        code = 'TSK-' + uuid.uuid4().hex[:6].upper()
        stamp = NOW()
        prompt_text = str(body.get('outcome_prompt') or '').strip() or f'Beneficiaries achieving positive outcome for {title}'
        payload = {
            'titles': {'en': title, 'rw': title, 'fr': title, 'sw': title},
            'outcome_prompt': {'en': prompt_text, 'rw': prompt_text, 'fr': prompt_text, 'sw': prompt_text},
            'unit_label': str(body.get('unit_label') or 'participants').strip(),
        }
        with self.tx():
            self.db.execute(
                'INSERT INTO v1_tasks VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (
                    tid, code, 'LocalRoot01', parish_code, goal['axis_code'],
                    goal['outcome_code'], goal['intervention_code'], goal['indicator_code'],
                    period['ap'], period['quarter'], due_date, planned_target,
                    min_threshold, target_threshold, frequency,
                    1 if body.get('is_sensitive') else 0, 'PENDING', user['id'],
                    self.enc(payload), stamp, stamp,
                ),
            )
            self.audit(user, 'V1_CREATE_TASK', tid, {'code': code, 'outcome': goal['outcome_code']})
            row = self.db.execute('SELECT * FROM v1_tasks WHERE id=?', (tid,)).fetchone()
            return self._serial_task(row, [])

    def v1_submit_log(self, user, body):
        from engine import Fault, NOW
        task_id = str(body.get('task_id', '')).strip()
        with self.tx():
            task_row = self.db.execute('SELECT * FROM v1_tasks WHERE id=?', (task_id,)).fetchone()
            if not task_row:
                raise Fault(404, 'plain_error', 'Selected task was not found. Please refresh your task list.')

            event_date = str(body.get('event_date', '')).strip()
            period = compute_pastoral(event_date)
            if not period:
                raise Fault(422, 'plain_error', 'Pick a date between 1 Sept 2026 and 31 Aug 2031.')

            parish_code = str(body.get('parish_code') or task_row['parish_code']).strip()
            centrale = str(body.get('centrale') or '').strip()
            if not centrale:
                raise Fault(422, 'plain_error', 'Select a Centrale or community location in Step 1.')

            collection_mode = str(body.get('collection_mode') or 'ENUMERATOR').strip()
            if collection_mode not in ('ENUMERATOR', 'FOCUS_GROUP', 'SURVEY', 'OBSERVATION'):
                collection_mode = 'ENUMERATOR'

            try:
                women = max(0, int(body.get('women', 0)))
                men = max(0, int(body.get('men', 0)))
                total = women + men  # Auto-calculated so addition errors never happen
                age_u18 = max(0, int(body.get('age_u18', 0)))
                age_18_35 = max(0, int(body.get('age_18_35', 0)))
                age_36_64 = max(0, int(body.get('age_36_64', 0)))
                age_65 = max(0, int(body.get('age_65', 0)))
                disability = max(0, int(body.get('disability', 0)))
                sessions_done = max(0.0, float(body.get('sessions_done', total or 1)))
                outcome_value = max(0.0, float(body.get('outcome_value', 0)))
            except (ValueError, TypeError):
                raise Fault(422, 'plain_error', 'Enter valid whole numbers for participant counts.')

            age_sum = age_u18 + age_18_35 + age_36_64 + age_65
            if total > 0 and age_sum != total:
                diff = total - age_sum
                action = f'Add {diff} more.' if diff > 0 else f'Remove {abs(diff)}.'
                raise Fault(
                    422,
                    'plain_error',
                    f'Age group numbers add up to {age_sum}, but you entered {total} total participants. {action}',
                )

            if disability > total:
                raise Fault(
                    422,
                    'plain_error',
                    f'Participants with disabilities ({disability}) cannot be higher than total participants ({total}).',
                )

            issue_severity = str(body.get('issue_severity') or 'NONE').upper()
            if issue_severity not in ('NONE', 'MINOR', 'MEDIUM', 'MAJOR'):
                issue_severity = 'NONE'
            raw_cats = body.get('issue_categories') or []
            if isinstance(raw_cats, str):
                raw_cats = [x.strip() for x in raw_cats.split(',') if x.strip()]
            issue_categories = ','.join(c for c in raw_cats if c in ISSUE_CATEGORIES)
            issue_note = str(body.get('issue_note') or '').strip()

            min_threshold = float(task_row['min_threshold'] or 0)
            target_threshold = float(task_row['target_threshold'] or 0)
            if outcome_value < min_threshold and (issue_severity == 'NONE' or not issue_note):
                raise Fault(
                    422,
                    'plain_error',
                    f'Target not reached ({int(outcome_value)} of {int(min_threshold)}). Please select an issue level and add a short note.',
                )

            if task_row['is_sensitive'] and not body.get('privacy_confirmed'):
                raise Fault(
                    422,
                    'plain_error',
                    'Please check the box confirming no victim names or private medical details are in the notes.',
                )

            impact_rating = compute_impact_rating(outcome_value, min_threshold, target_threshold)
            stamp = NOW()
            existing_id = body.get('id')
            parent_id = None
            version = 1
            revision = 1
            reviews = []

            if existing_id:
                old = self.db.execute('SELECT * FROM v1_logs WHERE id=?', (existing_id,)).fetchone()
                if old:
                    old_payload = self.dec(old['body'])
                    reviews = old_payload.get('reviews', [])
                    if old['stage'] == 'DIOCESE_APPROVED':
                        # Silent automatic amendment creation
                        parent_id = old['id']
                        version = old['version'] + 1
                        existing_id = None
                    else:
                        revision = old['revision'] + 1
                        version = old['version']

            payload = {
                'duration_hours': body.get('duration_hours', 2),
                'consensus_notes': str(body.get('consensus_notes', '')).strip(),
                'respondents': body.get('respondents', []),
                'forms_distributed': body.get('forms_distributed', 0),
                'forms_returned': body.get('forms_returned', 0),
                'checklist': body.get('checklist', []),
                'observation_notes': str(body.get('observation_notes', '')).strip(),
                'story_of_change': str(body.get('story_of_change', '')).strip(),
                'attachments': body.get('attachments', []),
                'privacy_confirmed': bool(body.get('privacy_confirmed', True)),
                'reviews': reviews,
                'correction_step': None,
                'correction_note': '',
            }

            if existing_id:
                self.db.execute(
                    '''UPDATE v1_logs SET
                        parish_code=?, centrale=?, event_date=?, ap=?, pastoral_quarter=?,
                        collection_mode=?, women=?, men=?, total=?,
                        age_u18=?, age_18_35=?, age_36_64=?, age_65=?, disability=?,
                        sessions_done=?, outcome_value=?, impact_rating=?,
                        issue_severity=?, issue_categories=?, issue_note=?,
                        stage='SUBMITTED', revision=?, body=?, updated=?
                       WHERE id=?''',
                    (
                        parish_code, centrale, event_date, period['ap'], period['quarter'],
                        collection_mode, women, men, total,
                        age_u18, age_18_35, age_36_64, age_65, disability,
                        sessions_done, outcome_value, impact_rating,
                        issue_severity, issue_categories, issue_note,
                        revision, self.enc(payload), stamp, existing_id,
                    ),
                )
                lid = existing_id
            else:
                lid = 'log-' + uuid.uuid4().hex[:12]
                self.db.execute(
                    'INSERT INTO v1_logs VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (
                        lid, task_id, 'LocalRoot01', parish_code, centrale,
                        event_date, period['ap'], period['quarter'], collection_mode,
                        women, men, total, age_u18, age_18_35, age_36_64, age_65,
                        disability, sessions_done, outcome_value, impact_rating,
                        issue_severity, issue_categories, issue_note, 'SUBMITTED',
                        revision, version, parent_id, user['id'], self.enc(payload),
                        stamp, stamp,
                    ),
                )

            self.audit(user, 'V1_SUBMIT_LOG', lid, {'task_id': task_id, 'impact': impact_rating, 'total': total})
            saved = self.db.execute('SELECT * FROM v1_logs WHERE id=?', (lid,)).fetchone()
            return self._serial_log(saved)

    def v1_review_log(self, user, body):
        from engine import Fault, NOW
        log_id = str(body.get('id', '')).strip()
        target_stage = str(body.get('target_stage', '')).strip()
        note = str(body.get('note', '')).strip()
        correction_step = body.get('correction_step')
        if target_stage not in ('PARISH_APPROVED', 'DIOCESE_APPROVED', 'CORRECTION'):
            raise Fault(422, 'plain_error', 'Choose a valid review action (Parish Approve, Diocese Approve, or Return).')
        if target_stage == 'CORRECTION' and not note:
            raise Fault(422, 'plain_error', 'Please write a short note explaining what needs to be fixed.')

        with self.tx():
            row = self.db.execute('SELECT * FROM v1_logs WHERE id=?', (log_id,)).fetchone()
            if not row:
                raise Fault(404, 'plain_error', 'Activity report was not found.')
            payload = self.dec(row['body'])
            reviews = payload.get('reviews', [])
            stamp = NOW()
            reviews.append({
                'stage': target_stage,
                'actor': user.get('name') or user['id'],
                'time': stamp,
                'note': note or ('Verified at Parish level' if target_stage == 'PARISH_APPROVED' else 'Finalized by Diocese'),
                'step': correction_step,
            })
            payload['reviews'] = reviews
            payload['correction_step'] = int(correction_step) if target_stage == 'CORRECTION' and correction_step else None
            payload['correction_note'] = note if target_stage == 'CORRECTION' else ''

            self.db.execute(
                'UPDATE v1_logs SET stage=?, revision=revision+1, body=?, updated=? WHERE id=?',
                (target_stage, self.enc(payload), stamp, log_id),
            )
            self.audit(user, 'V1_REVIEW_LOG', log_id, {'stage': target_stage})
            updated = self.db.execute('SELECT * FROM v1_logs WHERE id=?', (log_id,)).fetchone()
            return self._serial_log(updated)

    def _compute_v1_analytics(self, tasks, logs, parishes, parish_filter=None, period_filter=None, axis_filter=None):
        active_logs = [l for l in logs if l.get('parent_id') is None]
        if parish_filter:
            tasks = [t for t in tasks if t['parish_code'] == parish_filter]
            active_logs = [l for l in active_logs if l['parish_code'] == parish_filter]
        if period_filter and period_filter in ('Q1', 'Q2', 'Q3', 'Q4'):
            active_logs = [l for l in active_logs if l['pastoral_quarter'] == period_filter]
        elif period_filter and period_filter.startswith('AP'):
            tasks = [t for t in tasks if t['ap'] == period_filter]
            active_logs = [l for l in active_logs if l['ap'] == period_filter]
        if axis_filter:
            tasks = [t for t in tasks if t['axis'] == axis_filter]
            active_logs = [l for l in active_logs if l.get('axis') == axis_filter]

        total_women = sum(l['women'] for l in active_logs)
        total_men = sum(l['men'] for l in active_logs)
        total_participants = sum(l['total'] for l in active_logs)
        total_disability = sum(l['disability'] for l in active_logs)
        avg_completion = round(sum(t['completion_pct'] for t in tasks) / len(tasks), 1) if tasks else 0.0

        impact_counts = {'GREAT': 0, 'GOOD': 0, 'LOW': 0, 'PENDING': 0}
        for t in tasks:
            impact_counts[t['impact_rating']] = impact_counts.get(t['impact_rating'], 0) + 1

        # Issue & Bottleneck Aggregation
        issue_logs = [l for l in active_logs if l['issue_severity'] in ('MINOR', 'MEDIUM', 'MAJOR')]
        cat_counts = {c: 0 for c in ISSUE_CATEGORIES}
        for l in issue_logs:
            for c in l['issue_categories']:
                if c in cat_counts:
                    cat_counts[c] += 1
        total_cat_mentions = sum(cat_counts.values())
        top_cat = max(cat_counts.items(), key=lambda x: x[1])[0] if total_cat_mentions > 0 else None
        top_pct = round((cat_counts[top_cat] / len(issue_logs)) * 100) if top_cat and issue_logs else 0

        # Monthly Issue Trends series (Sept - Feb pastoral months)
        months_order = [
            ('2026-09', 'Sept'), ('2026-10', 'Oct'), ('2026-11', 'Nov'),
            ('2026-12', 'Dec'), ('2027-01', 'Jan'), ('2027-02', 'Feb'),
        ]
        monthly_trends = []
        for ym, mlabel in months_order:
            mlogs = [l for l in active_logs if l['event_date'].startswith(ym)]
            monthly_trends.append({
                'month': mlabel,
                'ym': ym,
                'MINOR': sum(1 for l in mlogs if l['issue_severity'] == 'MINOR'),
                'MEDIUM': sum(1 for l in mlogs if l['issue_severity'] == 'MEDIUM'),
                'MAJOR': sum(1 for l in mlogs if l['issue_severity'] == 'MAJOR'),
                'by_category': {
                    c: sum(1 for l in mlogs if c in l['issue_categories'] and l['issue_severity'] != 'NONE')
                    for c in ISSUE_CATEGORIES
                },
            })

        # Outcome Accordions
        goals = self._goals_catalogue()
        accordions = []
        for g in goals:
            if axis_filter and g['axis_code'] != axis_filter:
                continue
            otasks = [t for t in tasks if t['outcome'] == g['outcome_code']]
            ologs = [l for l in active_logs if l.get('outcome') == g['outcome_code']]
            if not otasks and not ologs:
                continue
            comp_pct = round(sum(t['completion_pct'] for t in otasks) / len(otasks), 1) if otasks else 0.0
            ratings = [t['impact_rating'] for t in otasks if t['impact_rating'] != 'PENDING']
            if 'LOW' in ratings:
                overall_impact = 'LOW'
            elif 'GREAT' in ratings and all(r == 'GREAT' for r in ratings):
                overall_impact = 'GREAT'
            elif ratings:
                overall_impact = 'GOOD'
            else:
                overall_impact = 'PENDING'
            documents = []
            stories = []
            issues = []
            for l in ologs:
                for att in l['payload'].get('attachments', []):
                    documents.append({'name': att, 'parish': l.get('parish_name', l['parish_code']), 'date': l['event_date']})
                if l['payload'].get('story_of_change'):
                    stories.append({'text': l['payload']['story_of_change'], 'parish': l.get('parish_name', l['parish_code']), 'date': l['event_date']})
                if l['issue_severity'] != 'NONE' and l['issue_note']:
                    issues.append({
                        'severity': l['issue_severity'],
                        'categories': l['issue_categories'],
                        'note': l['issue_note'],
                        'parish': l.get('parish_name', l['parish_code']),
                        'date': l['event_date'],
                    })
            accordions.append({
                'outcome_code': g['outcome_code'],
                'axis_code': g['axis_code'],
                'axis_labels': g['axis_labels'],
                'labels': g['labels'],
                'completion_pct': comp_pct,
                'impact_rating': overall_impact,
                'tasks': otasks,
                'documents': documents,
                'stories': stories,
                'issues': issues,
            })

        # Regional Parish Map nodes
        all_tasks_unfiltered = [self._serial_task(t, logs) for t in self.db.execute('SELECT * FROM v1_tasks')]
        parish_nodes = []
        for p in parishes:
            pcode = p['code']
            ptasks = [t for t in all_tasks_unfiltered if t['parish_code'] == pcode]
            plogs = [l for l in logs if l['parish_code'] == pcode and l.get('parent_id') is None]
            p_comp = round(sum(t['completion_pct'] for t in ptasks) / len(ptasks), 1) if ptasks else None
            major_issues = sum(1 for l in plogs if l['issue_severity'] == 'MAJOR')
            parish_nodes.append({
                'code': pcode,
                'name': p['name'],
                'district': p['district'],
                'x': p['x'],
                'y': p['y'],
                'tasks_count': len(ptasks),
                'logs_count': len(plogs),
                'completion_pct': p_comp,
                'major_issues': major_issues,
                'participants': sum(l['total'] for l in plogs),
                'status_color': 'red' if major_issues > 0 else ('green' if (p_comp or 0) >= 80 else ('amber' if p_comp is not None else 'neutral')),
            })

        return {
            'summary': {
                'tasks_total': len(tasks),
                'tasks_completed': sum(1 for t in tasks if t['completion_pct'] >= 100),
                'avg_completion_pct': avg_completion,
                'total_participants': total_participants,
                'women': total_women,
                'men': total_men,
                'disability': total_disability,
                'pending_approvals': sum(1 for l in active_logs if l['stage'] in ('SUBMITTED', 'PARISH_APPROVED')),
                'impact_counts': impact_counts,
                'top_bottleneck_category': top_cat,
                'top_bottleneck_pct': top_pct,
                'issues_total': len(issue_logs),
            },
            'category_breakdown': cat_counts,
            'monthly_trends': monthly_trends,
            'accordions': accordions,
            'parish_map': parish_nodes,
        }

    def v1_export(self, user, report_type='impact_summary', fmt='csv', preset='natural', lang='en', parish_filter=None):
        data = self.v1_bootstrap(user, parish_filter=parish_filter)
        logs = data['logs']
        tasks = data['tasks']

        headers_by_preset = {
            'natural': ['Date', 'Parish', 'Centrale', 'Activity Title', 'Collection Mode', 'Women', 'Men', 'Total Participants', 'Outcome Achieved', 'Impact Rating', 'Issue Level', 'Issue Categories', 'Field Notes', 'Approval Status'],
            'donor': ['Reporting_Date', 'Parish_Unit', 'Community_Site', 'Intervention_Activity', 'Methodology', 'Female_Beneficiaries', 'Male_Beneficiaries', 'Total_Reached', 'Outcome_Indicator_Actual', 'Outcome_Quality_Tier', 'Implementation_Risk_Level', 'Bottleneck_Tags', 'Qualitative_Variance_Note', 'Verification_Stage'],
            'dhis2': ['eventDate', 'orgUnitName', 'attributeOptionCombo', 'programStageName', 'collectionMethod', 'DE_WOMEN', 'DE_MEN', 'DE_TOTAL', 'DE_OUTCOME_VAL', 'DE_IMPACT_TIER', 'DE_ISSUE_SEV', 'DE_ISSUE_CAT', 'comment', 'status'],
        }
        headers = headers_by_preset.get(preset, headers_by_preset['natural'])
        stage_natural = {
            'SUBMITTED': 'Waiting for Parish Review',
            'PARISH_APPROVED': 'Approved by Parish',
            'DIOCESE_APPROVED': 'Approved by Diocese',
            'CORRECTION': 'Needs Correction',
        }

        rows = []
        for l in logs:
            if parish_filter and l['parish_code'] != parish_filter:
                continue
            title = l.get('task_titles', {}).get(lang) or l.get('task_titles', {}).get('en', 'Activity')
            rows.append([
                l['event_date'],
                l.get('parish_name', l['parish_code']),
                l['centrale'],
                title,
                l['collection_mode'],
                l['women'],
                l['men'],
                l['total'],
                l['outcome_value'],
                l['impact_rating'],
                l['issue_severity'],
                ', '.join(l['issue_categories']),
                l['issue_note'],
                stage_natural.get(l['stage'], l['stage']),
            ])

        if fmt == 'csv':
            out = io.StringIO()
            w = csv.writer(out)
            w.writerow(headers)
            for r in rows:
                w.writerow(r)
            return ('\ufeff' + out.getvalue()).encode('utf-8'), 'text/csv;charset=utf-8', f'APSE_{report_type}_{preset}.csv'

        if fmt == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = 'APSE Export'
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1E3A2F')
            for r in rows:
                ws.append(r)
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', f'APSE_{report_type}_{preset}.xlsx'

        from docx import Document
        doc = Document()
        doc.add_heading('CDJP Butare — Strategy & Impact Summary Report', 0)
        doc.add_paragraph(f'Generated on {date.today().isoformat()} · Total Activities Logged: {len(rows)}')
        table = doc.add_table(rows=1, cols=7)
        table.style = 'Table Grid'
        for idx, h in enumerate(['Date', 'Parish', 'Activity', 'Total', 'Outcome', 'Impact', 'Status']):
            table.rows[0].cells[idx].text = h
        for r in rows:
            row_cells = table.add_row().cells
            for idx, val in enumerate([r[0], r[1], r[3], str(r[7]), str(r[8]), r[9], r[13]]):
                row_cells[idx].text = str(val)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', f'APSE_{report_type}.docx'

    def v1_import_csv(self, user, body):
        from engine import Fault
        raw_text = str(body.get('csv_text', '')).strip()
        if not raw_text:
            raise Fault(422, 'plain_error', 'Choose a CSV file with at least one activity row.')
        reader = csv.reader(io.StringIO(raw_text))
        try:
            headers = [h.strip() for h in next(reader)]
        except StopIteration:
            raise Fault(422, 'plain_error', 'Your CSV file is empty.')
        data_rows = [r for r in reader if any(c.strip() for c in r)]
        mapping = body.get('mapping') or {}

        required_roles = [
            ('parish', 'Parish_Name', ['parish', 'parish_name', 'paroisse', 'paruwasi', 'unit', 'parish_unit']),
            ('centrale', 'Centrale', ['centrale', 'santarali', 'community_site', 'site']),
            ('event_date', 'Activity_Date', ['date', 'event_date', 'activity_date', 'reporting_date', 'eventdate', 'itariki']),
            ('women', 'Women_Count', ['women', 'women_count', 'female', 'female_beneficiaries', 'abagore', 'de_women']),
            ('men', 'Men_Count', ['men', 'men_count', 'male', 'male_beneficiaries', 'abagabo', 'de_men']),
            ('outcome_value', 'Outcome_Achieved', ['outcome', 'outcome_achieved', 'outcome_value', 'outcome_indicator_actual', 'de_outcome_val']),
        ]
        norm_headers = {h.lower().replace(' ', '_'): h for h in headers}
        auto_map = {}
        missing = []
        for key, friendly_col, aliases in required_roles:
            chosen = mapping.get(key)
            if chosen and chosen in headers:
                auto_map[key] = chosen
                continue
            found = next((norm_headers[a] for a in aliases if a in norm_headers), None)
            if found:
                auto_map[key] = found
            else:
                missing.append({'key': key, 'column': friendly_col, 'message': f"Column '{friendly_col}' is missing in your file. Select which column contains the {friendly_col.replace('_', ' ')}."})

        if body.get('commit') and missing:
            raise Fault(422, 'plain_error', missing[0]['message'])

        preview_rows = []
        parishes, assigned_parish, _ = self._parishes_and_centrales(user)
        parish_by_name = {p['name'].lower(): p['code'] for p in parishes}
        default_task = self.db.execute('SELECT id FROM v1_tasks LIMIT 1').fetchone()
        default_task_id = body.get('task_id') or (default_task['id'] if default_task else 'task-a2-gakoma-01')

        imported_count = 0
        for idx, r in enumerate(data_rows[:100], start=1):
            row_dict = {headers[i]: (r[i].strip() if i < len(r) else '') for i in range(len(headers))}
            if not missing:
                pname = row_dict.get(auto_map['parish'], '')
                pcode = parish_by_name.get(pname.lower(), assigned_parish)
                centrale = row_dict.get(auto_map['centrale'], 'Centre') or 'Centre'
                edate = row_dict.get(auto_map['event_date'], '2026-10-20') or '2026-10-20'
                try:
                    w = int(float(row_dict.get(auto_map['women'], 0) or 0))
                    m = int(float(row_dict.get(auto_map['men'], 0) or 0))
                    out_v = float(row_dict.get(auto_map['outcome_value'], w + m) or 0)
                except ValueError:
                    w, m, out_v = 0, 0, 0
                preview_rows.append({
                    'row': idx,
                    'parish': pname or pcode,
                    'centrale': centrale,
                    'event_date': edate,
                    'women': w,
                    'men': m,
                    'total': w + m,
                    'outcome_value': out_v,
                })
                if body.get('commit'):
                    ptask = self.db.execute('SELECT id, min_threshold FROM v1_tasks WHERE parish_code=? LIMIT 1', (pcode,)).fetchone()
                    tid = body.get('task_id') or (ptask['id'] if ptask else default_task_id)
                    min_t = float(ptask['min_threshold']) if ptask else 12.0
                    self.v1_submit_log(user, {
                        'task_id': tid,
                        'parish_code': pcode,
                        'centrale': centrale,
                        'event_date': edate,
                        'collection_mode': 'SURVEY',
                        'women': w,
                        'men': m,
                        'age_18_35': w + m,
                        'sessions_done': w + m,
                        'outcome_value': out_v,
                        'issue_severity': 'MINOR' if out_v < min_t else 'NONE',
                        'issue_categories': ['Other'] if out_v < min_t else [],
                        'issue_note': 'Partial batch imported from CSV spreadsheet.' if out_v < min_t else '',
                        'privacy_confirmed': True,
                    })
                    imported_count += 1

        return {
            'headers': headers,
            'mapping': auto_map,
            'missing': missing,
            'preview': preview_rows[:15],
            'total_rows': len(data_rows),
            'imported': imported_count,
        }
