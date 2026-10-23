"""Bounded CSV/XLSX/JSON transfer with server-side mapping and draft-only imports."""
import csv,io,json,re,unicodedata,zipfile,hashlib
from datetime import date,datetime
MAX_BYTES=8*1024*1024;MAX_ROWS=1000;MAX_COLUMNS=180
class Transfers:
 def norm(self,x):return ''.join(c for c in unicodedata.normalize('NFKD',str(x)).lower()if c.isalnum())
 def transfer_columns(self,form):
  from engine import Fault
  f=self.form_map().get(form)
  if not f:raise Fault(422,'unknown_form')
  return ['event_date']+[k for k in f['fields'] if k not in ['state','approved','approval_ref']]
 def parse_transfer(self,data,fmt):
  from engine import Fault
  if not data or len(data)>MAX_BYTES:raise Fault(413,'import_file_limit')
  try:
   if fmt=='json':
    rows=json.loads(data.decode('utf-8-sig'));rows=rows.get('records')if isinstance(rows,dict)else rows
    if not isinstance(rows,list) or any(not isinstance(x,dict) for x in rows):raise Fault(422,'import_json_rows_required')
    # Export rows deliberately retain workflow metadata; import only event_date/values.
    rows=[dict(event_date=x.get('event_date'),**x['values'])if isinstance(x.get('values'),dict)else x for x in rows]
    headers=list(dict.fromkeys(k for x in rows for k in x));table=[[x.get(k,'')for k in headers]for x in rows]
   elif fmt=='csv':
    raw=data.decode('utf-8-sig');dialect=csv.Sniffer().sniff(raw[:8000],delimiters=',;\t')if '\n'in raw else csv.excel
    reader=csv.reader(io.StringIO(raw),dialect);headers=next(reader);table=list(reader)
   elif fmt=='xlsx':
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(data))as z:
     if sum(x.file_size for x in z.infolist())>40*1024*1024 or len(z.infolist())>200:raise Fault(413,'xlsx_expansion_limit')
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False);sheet=wb.worksheets[0]
    if sheet.max_row>MAX_ROWS+1 or sheet.max_column>MAX_COLUMNS:raise Fault(413,'import_shape_limit')
    rows=[]
    for row in sheet.iter_rows():
     if any(c.data_type=='f'for c in row):raise Fault(422,'excel_formulas_not_imported')
     rows.append([c.value for c in row])
    wb.close();headers=rows[0];table=rows[1:]
   else:raise Fault(422,'unsupported_import_format')
  except Fault:raise
  except Exception:raise Fault(422,'invalid_import_file')
  if len(table)>MAX_ROWS or len(headers)>MAX_COLUMNS:raise Fault(413,'import_shape_limit')
  if not headers or any(x in [None,'']for x in headers) or len({self.norm(x)for x in headers})!=len(headers):raise Fault(422,'invalid_or_duplicate_headers')
  headers=[str(x).strip()for x in headers];out=[]
  for row in table:
   if not any(v not in [None,'']for v in row):continue
   if len(row)>len(headers) and any(x not in [None,'']for x in row[len(headers):]):raise Fault(422,'extra_import_cells')
   def cell(x):
    if isinstance(x,(date,datetime)):return x.date().isoformat()if isinstance(x,datetime)else x.isoformat()
    if isinstance(x,bool):return 'true'if x else 'false'
    if x is None:return ''
    if isinstance(x,(dict,list)):raise Fault(422,'nested_import_cell')
    return str(x).strip()
   out.append({h:cell(row[i])if i<len(row)else ''for i,h in enumerate(headers)})
  return headers,out
 def transfer_mapping(self,form,headers,mapping=None):
  from engine import Fault
  allowed=self.transfer_columns(form);fields=self.configuration()['fields'];aliases={}
  for key in allowed:
   for name in [key]+list(fields.get(key,{}).get('label',{}).values()):aliases.setdefault(self.norm(name),set()).add(key)
  for name in ['Date événement','Event date','Itariki y’igikorwa','Tarehe ya tukio']:aliases.setdefault(self.norm(name),set()).add('event_date')
  if mapping is None:mapping={h:next(iter(aliases[self.norm(h)]))for h in headers if len(aliases.get(self.norm(h),set()))==1}
  if not isinstance(mapping,dict) or set(mapping)-set(headers) or any(v not in allowed+['__IGNORE__'] for v in mapping.values()) or len([v for v in mapping.values()if v!='__IGNORE__'])!=len({v for v in mapping.values()if v!='__IGNORE__'}):raise Fault(422,'invalid_import_mapping')
  return mapping
 def prepare_transfer(self,user,form,unit,path,data,fmt,mapping=None):
  from engine import Fault
  self.check(user,dict(form=form,unit=unit,path=path,author=user['id']),True);headers,rows=self.parse_transfer(data,fmt);mapping=self.transfer_mapping(form,headers,mapping);out=[];seen=set()
  for n,row in enumerate(rows,2):
   values={target:row[source]for source,target in mapping.items()if target!='__IGNORE__' and row[source]!=''};event=values.pop('event_date','');error=None
   key=hashlib.sha256(__import__('engine').canonical([hashlib.sha256(data).hexdigest(),form,unit,mapping,n]).encode()).hexdigest()
   if not values.get('record'):values['record']=self.generated_record_code(user,form,unit,key)
   if 'capture_profile'in self.form_map()[form]['fields']:values['capture_profile']='V8'
   try:
    v=self.validate_plan(user,form,values,event,unit)
    if form in ['F31','F32'] and not self.roles(user)&{'APSE_HR','APSE_MEL'} and v.get('owner')!=user['id']:raise Fault(403,'own_staff_record_only')
    self._links(user,v,form)
    if v.get('record')in seen:raise Fault(409,'duplicate_record_in_import')
    key=hashlib.sha256(__import__('engine').canonical([hashlib.sha256(data).hexdigest(),form,unit,mapping,n]).encode()).hexdigest()
    old=self.db.execute('SELECT record_id FROM requests WHERE author=? AND key=?',(user['id'],key)).fetchone()
    for existing in self.db.execute('SELECT id,body FROM records WHERE form=? AND unit=? AND parent IS NULL',(form,unit)):
     if self.dec(existing['body']).get('record')==v.get('record') and (not old or old['record_id']!=existing['id']):raise Fault(409,'duplicate_business_code')
    seen.add(v.get('record'))
   except Fault as fault:error={'code':fault.code,'details':fault.details}
   out.append(dict(row=n,event_date=event,values=values,error=error))
  return dict(headers=headers,mapping=mapping,unmapped=[h for h in headers if h not in mapping],rows=out,valid=sum(x['error']is None for x in out),invalid=sum(x['error']is not None for x in out),fingerprint=hashlib.sha256(data).hexdigest(),writes=False)
 def import_transfer(self,user,form,unit,path,data,fmt,mapping=None):
  from engine import Fault,canonical
  result=self.prepare_transfer(user,form,unit,path,data,fmt,mapping)
  if result['invalid']or result['unmapped']:raise Fault(422,'import_preview_requires_correction',{'invalidRows':result['invalid'],'unmapped':result['unmapped']})
  rows=[]
  for row in result['rows']:
   key=hashlib.sha256(canonical([result['fingerprint'],form,unit,result['mapping'],row['row']]).encode()).hexdigest()
   try:
    record=self.create(user,form,unit,path,row['event_date'],row['values'],key);rows.append(dict(row=row['row'],id=record['id'],state=record['state']))
   except Fault as fault:rows.append(dict(row=row['row'],error=fault.code))
  with self.tx():self.audit(user,'IMPORT_DRAFTS',form,{'fileHash':result['fingerprint'],'rows':len(rows),'errors':sum('error'in x for x in rows)})
  return dict(rows=rows,createdOrReused=sum('id'in x for x in rows),errors=sum('error'in x for x in rows),atomic=False,approval='DRAFT_ONLY')
 def export_transfer(self,user,form,start,end,unit,fmt,include_sensitive=False,approved=True,language='fr',template=False):
  from engine import Fault
  from reports import safe
  if fmt not in ['csv','xlsx','json']or language not in ['fr','en','rw','sw']:raise Fault(422,'invalid_export_options')
  if template:
   columns=self.transfer_columns(form);rows=[]
  else:
   if form and form not in self.form_map():raise Fault(422,'unknown_form')
   if include_sensitive and not self.roles(user)&{'APSE_MEL','APSE_SENSITIVE','APSE_FINANCE','APSE_HR'}:raise Fault(403,'sensitive_export_role_required')
   rows=[r for r in self.listing(user,start,end,form=form,approved=approved) if (not unit or r['unit']==unit or unit in r['path'].split('/')) and (include_sensitive or not self.form_map()[r['form']]['sensitive'])]
   columns=['id','form','unit','event_date','state','revision']+list(dict.fromkeys(k for r in rows for k in r['values']))
  if fmt=='json':data=json.dumps({'columns':columns,'records':rows,'template':template,'workflowImport':'DRAFT_ONLY'},ensure_ascii=False,indent=2).encode();mime='application/json'
  else:
   table=[]
   for r in rows:table.append([r.get(k,r['values'].get(k,''))if k in ['id','form','unit','event_date','state','revision']else r['values'].get(k,'')for k in columns])
   if fmt=='csv':
    out=io.StringIO();writer=csv.writer(out);writer.writerow(columns)
    for row in table:writer.writerow([safe(x)for x in row])
    data=('\ufeff'+out.getvalue()).encode();mime='text/csv;charset=utf-8'
   else:
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();sheet=wb.active;sheet.title='APSE';sheet.append(columns)
    for row in table:sheet.append([safe(x)if isinstance(x,str)else x for x in row])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(name='Arial',bold=True);cell.fill=PatternFill('solid',fgColor='FDF751')
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=24
    if template:
     guide=wb.create_sheet('Guide');guide.append(['code','label','required','type'])
     for key in columns:guide.append([key,self.configuration()['fields'].get(key,{}).get('label',{}).get(language,'Date'),key in self.form_map()[form]['required']or key=='event_date',self.configuration()['fields'].get(key,{}).get('type','DATE')])
    out=io.BytesIO();wb.save(out);data=out.getvalue();mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
  with self.tx():self.audit(user,'EXPORT_DATA'if not template else 'EXPORT_TEMPLATE',form or'ALL',{'format':fmt,'rows':len(rows),'sensitiveExplicit':include_sensitive})
  return data,mime
