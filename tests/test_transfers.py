import unittest,sys,tempfile,json,io,csv
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'service'))
from engine import Engine,Fault
from cryptography.fernet import Fernet
from openpyxl import Workbook,load_workbook
class TransferTests(unittest.TestCase):
 def setUp(self):
  self.e=Engine(':memory:',Fernet.generate_key(),b't'*32);self.user=dict(id='STAFF',roles=['APSE_STAFF'],paths=['/ROOT']);self.mel=dict(id='MEL',roles=['APSE_MEL'],paths=['/ROOT'])
 def tearDown(self):self.e.db.close()
 def data(self,**extra):
  r=dict(event_date='2026-10-20',record='IMPORT-1',axis='A1',result='A1.R1',activity='ACT',title='Training',owner='STAFF',total=10,women=6,men=4,undisclosed=0,work_type='OPERATIONAL',reason='Documented unplanned activity');r.update(extra);return r
 def raw(self,rows=None):return json.dumps(rows or[self.data()]).encode()
 def preview(self,data=None,fmt='json',mapping=None):return self.e.prepare_transfer(self.user,'F07','P1','/ROOT/P1',data or self.raw(),fmt,mapping)
 def commit(self,data=None):return self.e.import_transfer(self.user,'F07','P1','/ROOT/P1',data or self.raw(),'json')
 def fault(self,code,fn):
  with self.assertRaises(Fault)as cm:fn()
  self.assertEqual(cm.exception.code,code)
 def test_preview_does_not_write(self):
  p=self.preview();self.assertEqual(p['valid'],1);self.assertEqual(self.e.db.execute('SELECT COUNT(*)FROM records').fetchone()[0],0)
 def test_import_is_draft_only_and_idempotent(self):
  a=self.commit();b=self.commit();self.assertEqual(a['rows'][0]['id'],b['rows'][0]['id']);self.assertEqual(a['rows'][0]['state'],'DRAFT');self.assertEqual(self.e.db.execute('SELECT COUNT(*)FROM records').fetchone()[0],1)
 def test_validation_before_writes(self):
  p=self.preview(self.raw([self.data(),self.data(record='SECOND',axis='A2') ]));self.assertEqual(p['invalid'],1);self.fault('import_preview_requires_correction',lambda:self.commit(self.raw([self.data(),self.data(record='SECOND',axis='A2')])));self.assertEqual(self.e.db.execute('SELECT COUNT(*)FROM records').fetchone()[0],0)
 def test_duplicate_codes_in_file(self):self.assertEqual(self.preview(self.raw([self.data(),self.data()]))['rows'][1]['error']['code'],'duplicate_record_in_import')
 def test_existing_code_cannot_be_reimported_from_different_file(self):
  self.commit();self.assertEqual(self.preview(self.raw([self.data(title='Changed')]))['rows'][0]['error']['code'],'duplicate_business_code')
 def test_scope_enforced_before_import(self):self.fault('access_denied',lambda:self.e.prepare_transfer(dict(self.user,paths=['/OTHER']),'F07','P1','/ROOT/P1',self.raw(),'json'))
 def test_ccjp_cannot_import_plan(self):self.fault('access_denied',lambda:self.e.prepare_transfer(dict(self.user,roles=['APSE_CCJP']),'F03','P1','/ROOT/P1',self.raw(),'json'))
 def test_reserved_workflow_cannot_be_mapped(self):self.fault('invalid_import_mapping',lambda:self.preview(mapping={'state':'state'}))
 def test_explicit_ignore_columns(self):
  raw=self.raw([self.data(note='not a target field')]);p=self.preview(raw);self.assertEqual(p['unmapped'],['note']);m=p['mapping'];m['note']='__IGNORE__';self.assertEqual(self.preview(raw,mapping=m)['unmapped'],[])
 def test_french_column_auto_mapping(self):
  row=self.data();row['Code du dossier']=row.pop('record');self.assertEqual(self.preview(self.raw([row]))['mapping']['Code du dossier'],'record')
 def test_csv_semicolon_and_bom(self):
  r=self.data();s=io.StringIO();w=csv.DictWriter(s,fieldnames=r,delimiter=';');w.writeheader();w.writerow(r);self.assertEqual(self.preview(('\ufeff'+s.getvalue()).encode(),'csv')['valid'],1)
 def test_xlsx_and_date_cells(self):
  import datetime
  r=self.data();r['event_date']=datetime.datetime(2026,10,20);w=Workbook();w.active.append(list(r));w.active.append(list(r.values()));out=io.BytesIO();w.save(out);self.assertEqual(self.preview(out.getvalue(),'xlsx')['rows'][0]['event_date'],'2026-10-20')
 def test_excel_formulas_rejected(self):
  w=Workbook();w.active.append(['event_date','title']);w.active.append(['2026-10-20','=1+1']);out=io.BytesIO();w.save(out);self.fault('excel_formulas_not_imported',lambda:self.preview(out.getvalue(),'xlsx'))
 def test_all_export_formats_and_templates(self):
  self.commit()
  for fmt in ['csv','xlsx','json']:
   b,m=self.e.export_transfer(self.user,'F07','2026-01-01','2027-01-01','ROOT',fmt,approved=False);self.assertTrue(b)
   if fmt=='json':self.assertEqual(len(json.loads(b)['records']),1)
   b,m=self.e.export_transfer(self.user,'F07','2026-01-01','2027-01-01','ROOT',fmt,template=True);self.assertTrue(b)
   if fmt=='xlsx':self.assertIn('Guide',load_workbook(io.BytesIO(b)).sheetnames)
 def test_sensitive_export_requires_explicit_option_and_role(self):
  self.e.create(self.mel,'F09','P1','/ROOT/P1','2026-10-20',dict(record='SECRET',case='SECRET-CODE',consent='true'),'SECRET_KEY_123456789');self.fault('sensitive_export_role_required',lambda:self.e.export_transfer(self.user,None,'2026-01-01','2027-01-01',None,'json',include_sensitive=True,approved=False));b,m=self.e.export_transfer(self.mel,None,'2026-01-01','2027-01-01',None,'json',approved=False);self.assertNotIn(b'SECRET-CODE',b)
 def test_json_export_can_reimport_values_without_state(self):
  self.commit();b,m=self.e.export_transfer(self.user,'F07','2026-01-01','2027-01-01','ROOT','json',approved=False);headers,rows=self.e.parse_transfer(b,'json');self.assertNotIn('state',headers);self.assertIn('record',rows[0])
