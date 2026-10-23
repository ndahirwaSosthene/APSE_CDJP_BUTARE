import unittest,sys,io,json,copy,zipfile,tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'service'))
from engine import CONFIG,Engine,Fault
from rbm import model,render,NARRATIVE
from cryptography.fernet import Fernet
from openpyxl import load_workbook
class ResultsReportingTests(unittest.TestCase):
 def row(self,id,form,v,state='APPROVED',unit='P1'):
  return dict(id=id,form=form,values=v,state=state,unit=unit,event_date='2026-10-20',revision=2,completion=None)
 def data(self):
  return [self.row('OBS1','F42',dict(indicator='A1-K1',axis='A1',result='A1.R1',numerator='8',denominator='10',baseline='50',target='90',method='Paired follow-up',evidence='SAFE-CODE'))]
 def test_all_ps_results_including_missing_data(self):
  d=model([],CONFIG,'2026-09-01','2027-08-31','fr');self.assertEqual(len(d['sections']),5);self.assertEqual(sum(len(a['results'])for a in d['sections']),13);self.assertFalse(d['sections'][0]['results'][0]['observations'])
 def test_baseline_target_actual_and_variance(self):
  d=model(self.data(),CONFIG,'2026-09-01','2027-08-31','fr');o=d['sections'][0]['results'][0]['observations'][0];self.assertEqual(o['actual'],80);self.assertEqual(o['change'],30);self.assertEqual(o['variance'],-10)
 def test_units_never_averaged(self):
  rows=self.data()+[self.row('OBS2','F42',dict(indicator='A1-K1',axis='A1',result='A1.R1',numerator='1',denominator='10',method='Different scope',evidence='SAFE'),unit='P2')];d=model(rows,CONFIG,'2026-09-01','2027-08-31','fr');self.assertEqual([x['actual']for x in d['sections'][0]['results'][0]['observations']],[80,10])
 def test_no_invented_result_mapping(self):
  rows=self.data();del rows[0]['values']['result'];d=model(rows,CONFIG,'2026-09-01','2027-08-31','fr');self.assertEqual(len(d['unmapped']),1);self.assertFalse(d['sections'][0]['results'][0]['observations'])
 def test_zero_denominator_is_missing(self):
  rows=self.data();rows[0]['values']['denominator']='0';d=model(rows,CONFIG,'2026-09-01','2027-08-31','fr');self.assertIsNone(d['sections'][0]['results'][0]['observations'][0]['actual'])
 def test_drafts_and_sensitive_cases_do_not_enter_matrix(self):
  rows=self.data();rows[0]['state']='DRAFT';rows.append(self.row('PRIVATE','F09',dict(case='SECRET-IDENTITY',finding='SECRET')));d=model(rows,CONFIG,'2026-09-01','2027-08-31','fr');self.assertNotIn('SECRET',json.dumps(d));self.assertFalse(d['sections'][0]['results'][0]['observations'])
 def test_word_excel_csv_in_four_languages(self):
  for lang in ['fr','en','rw','sw']:
   d=model(self.data(),CONFIG,'2026-09-01','2027-08-31',lang,{k:'Human analysis'for k in NARRATIVE})
   for fmt in ['docx','xlsx','csv']:
    b,mime=render(d,lang,fmt);self.assertTrue(b)
    if fmt in ['docx','xlsx']:self.assertTrue(zipfile.is_zipfile(io.BytesIO(b)))
    if fmt=='xlsx':self.assertIn('Narrative',load_workbook(io.BytesIO(b)).sheetnames)
 def test_freeze_report_and_coordinator_can_open_snapshot(self):
  with tempfile.TemporaryDirectory()as tmp:
   e=Engine(tmp+'/db',Fernet.generate_key(),b'a'*32);mel=dict(id='MEL',roles=['APSE_MEL'],paths=['/ROOT']);coord=dict(id='COORD',roles=['APSE_COORDINATION'],paths=['/ROOT']);d=e.freeze_report(mel,'ROOT','/ROOT','R13','fr','docx','2026-09-01','2027-08-31',{k:'Reviewed narrative'for k in NARRATIVE});snap=e.report_snapshot(coord,d['snapshot']);self.assertTrue(snap['data'].startswith(b'PK'))
   with self.assertRaises(Fault) as cm:e.edit(mel,d['record']['id'],1,d['record']['values'],'2026-10-20')
   self.assertEqual(cm.exception.code,'results_report_requires_new_snapshot');e.db.close()
 def test_parish_and_diocese_scope_rollup(self):
  with tempfile.TemporaryDirectory()as tmp:
   e=Engine(tmp+'/db',Fernet.generate_key(),b'b'*32);u=dict(id='MEL',roles=['APSE_MEL'],paths=['/ROOT']);v=dict(record='OBS',indicator='A1-K1',axis='A1',result='A1.R1',numerator='8',denominator='10',method='Paired measure',evidence='SAFE');r=e.create(u,'F42','CEB1','/ROOT/P1/CEB1','2026-10-20',v,'SCOPE_KEY_1234567890');e.db.execute("UPDATE records SET state='APPROVED' WHERE id=?",(r['id'],));self.assertEqual(len(e.results_report_rows(u,'2026-09-01','2027-08-31','ROOT')),1);self.assertEqual(len(e.results_report_rows(u,'2026-09-01','2027-08-31','P1')),1);self.assertEqual(e.results_report_rows(u,'2026-09-01','2027-08-31','P2'),[]);e.db.close()
 def test_empty_narrative_cannot_be_approved(self):
  with tempfile.TemporaryDirectory()as tmp:
   e=Engine(tmp+'/db',Fernet.generate_key(),b'c'*32);mel=dict(id='MEL',roles=['APSE_MEL'],paths=['/ROOT']);coord=dict(id='COORD',roles=['APSE_COORDINATION'],paths=['/ROOT']);director=dict(id='DIR',roles=['APSE_DIRECTION'],paths=['/ROOT']);d=e.freeze_report(mel,'ROOT','/ROOT','R13','fr','docx','2026-09-01','2027-08-31',{});r=d['record']
   for target,u in [('SUBMITTED',mel),('MEL_VALIDATED',mel),('VERIFIED',coord)]:r=e.transition(u,r['id'],r['revision'],target,'Review')
   with self.assertRaises(Fault)as cm:e.transition(director,r['id'],r['revision'],'APPROVED','Approval')
   self.assertEqual(cm.exception.code,'results_narrative_required');e.db.close()
 def test_mismatched_kpi_axis_not_mapped_to_wrong_result(self):
  rows=self.data();rows[0]['values']['result']='A2.R1';rows[0]['values']['axis']='A2';d=model(rows,CONFIG,'2026-09-01','2027-08-31','fr');self.assertEqual(len(d['unmapped']),1)
